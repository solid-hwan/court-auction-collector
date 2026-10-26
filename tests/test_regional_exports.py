import csv,json,tempfile,unittest,hashlib
from pathlib import Path
from openpyxl import load_workbook
from court_auction_mcp.export import export_records,restore_excel,restore_csv
from court_auction_mcp.models import map_item
from fixtures import item

class RegionalExportTests(unittest.TestCase):
    def test_region_folders_and_spreadsheet_roundtrip(self):
        records=[map_item(item(mae=1,daepyoSidoCd='11',daepyoSiguCd='680',hjguSido='서울특별시',hjguSigu='강남구',bgPlaceRdAllAddr='서울특별시 강남구 가상로 1')),map_item(item(mae=2,daepyoSidoCd='26',daepyoSiguCd='350',hjguSido='부산광역시',hjguSigu='해운대구',dspslUsgNm='=1+2',bgPlaceRdAllAddr='부산광역시 해운대구 가상로 2')),map_item(item(mae=3))]
        with tempfile.TemporaryDirectory() as root:
            dest=Path(root)/'export'; result=export_records(iter(records),dest,{'partial':True,'saved_count':3},by_region=True)
            self.assertEqual(result['region_groups'],3)
            manifest=json.loads((dest/'manifest.json').read_text())
            grouped=[]
            for region in manifest['regions']:
                folder=dest/region['path']
                rows=list(restore_excel(folder/'records.xlsx')); self.assertEqual(rows,list(restore_csv(folder/'records.csv')))
                self.assertEqual(len(rows),region['rows']); grouped+=rows
            self.assertEqual({r['item_key'] for r in grouped},{r['item_key'] for r in records})
            for name,sha in manifest['sha256'].items(): self.assertEqual(hashlib.sha256((dest/name).read_bytes()).hexdigest(),sha)
            wb=load_workbook(dest/'records.xlsx'); ws=wb['Overview']
            self.assertEqual(ws.freeze_panes,'D2'); self.assertEqual(ws.auto_filter.ref,'A1:N4')
            self.assertEqual(ws['J2'].data_type,'n'); self.assertEqual(ws['H3'].data_type,'s')
            self.assertIn('강남구',ws['C2'].value); wb.close()
            with (dest/'overview.csv').open(encoding='utf-8-sig') as f: rows=list(csv.DictReader(f))
            self.assertEqual(rows[1]['용도'],"'=1+2")
            self.assertTrue((dest/'지역별'/'지역목록.csv').exists())

    def test_unknown_and_malicious_location_never_escapes(self):
        from court_auction_mcp.regions import folder_parts
        r=map_item(item(hjguSigu='../../malicious\\escape',hjguSido='서울특별시'))
        for part in folder_parts(r):
            self.assertNotIn('/',part); self.assertNotIn('\\',part); self.assertNotIn('..',part)
