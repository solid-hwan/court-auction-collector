"""허용된 원문만 내보내며 Excel의 셀 길이·수식 제한을 처리한다."""
import csv
import hashlib
import json
import os
import re
import shutil
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from .models import AuctionError, canonical


def text_row(sheet, values):
    cells=[]
    for value in values:
        cell=WriteOnlyCell(sheet, value=str(value))
        cell.data_type='s'
        cells.append(cell)
    sheet.append(cells)


def export_records(records, destination, context):
    destination=Path(destination)
    if destination.exists():
        raise AuctionError('export_destination_exists')
    destination.parent.mkdir(parents=True,exist_ok=True)
    staging=Path(tempfile.mkdtemp(prefix='.auction-export-',dir=destination.parent))
    try:
        workbook=Workbook(write_only=True)
        overview=workbook.create_sheet('Overview')
        sheet=workbook.create_sheet('records')
        overflow=workbook.create_sheet('long_values')
        metadata=workbook.create_sheet('Collection')
        text_row(overview,['법원','사건','매물','목적물','감정가','최저매각가','매각기일','유찰횟수'])
        text_row(sheet,['record_index','record_json'])
        text_row(overflow,['record_index','chunk_index','text'])
        count=0
        with (staging/'records.jsonl').open('w',encoding='utf-8') as jf, (staging/'records.csv').open('w',encoding='utf-8-sig',newline='') as cf:
            writer=csv.writer(cf)
            writer.writerow(['record_index','record_json'])
            for count, record in enumerate(records,1):
                value=canonical(record)
                jf.write(value+'\n')
                writer.writerow([count,value])
                n=record['normalized']
                shown=[n.get(k) for k in ('court_code','case_no','sale_item_no','object_no','appraisal_price','minimum_price','sale_date','fail_count')]
                # 개요는 편의용. 값·빈값 이유의 정본은 records 시트와 JSONL이다.
                text_row(overview,['' if v is None else v for v in shown])
                if len(value)>30000:
                    for i,start in enumerate(range(0,len(value),30000)):
                        text_row(overflow,[count,i,value[start:start+30000]])
                    text_row(sheet,[count,f'[long_values:{count}]'])
                else:
                    text_row(sheet,[count,value])
        safe_context={k:context.get(k) for k in ('job_id','status','reason','partial','query','calls','max_calls','next_page','saved_count','read_count','total','total_reason','total_changed','snapshot_guarantee')}
        # 쿼리는 한정된 코드/숫자/날짜만 허용되어 셀 한도를 초과하지 않는다.
        for k,v in safe_context.items():
            text_row(metadata,[k,canonical(v)])
        workbook.save(staging/'records.xlsx')
        manifest={'rows':count,'collection':safe_context,'canonical':'records.jsonl','csv':'record_json is a JSON object; parse with json.loads','excel':'restore_excel(path) joins ordered long_values chunks','sha256':{name:hashlib.sha256((staging/name).read_bytes()).hexdigest() for name in ('records.jsonl','records.csv','records.xlsx')}}
        (staging/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        # 디렉터리 교체는 기존 파일 덮어쓰기를 허용하지 않는다.
        if destination.exists():
            raise AuctionError('export_destination_exists')
        os.rename(staging,destination)
        return {'ok':True,'directory':str(destination.resolve()),'rows':count,'partial':safe_context['partial'],'files':list(manifest['sha256'])+['manifest.json']}
    finally:
        if staging.exists():
            shutil.rmtree(staging)


def restore_excel(path):
    workbook=load_workbook(path,read_only=True,data_only=False)
    try:
        chunks={}
        for index,order,text in workbook['long_values'].iter_rows(min_row=2,values_only=True):
            chunks.setdefault(index,[]).append((int(order),text))
        for index,value in workbook['records'].iter_rows(min_row=2,values_only=True):
            if value==f'[long_values:{index}]':
                value=''.join(t for _,t in sorted(chunks[index]))
            yield json.loads(value)
    finally:
        workbook.close()


def restore_csv(path):
    """Python csv 기본 128KiB 제한보다 긴 원문도 복원한다."""
    previous=csv.field_size_limit()
    try:
        csv.field_size_limit(256*1024*1024)
        with Path(path).open(encoding='utf-8-sig',newline='') as stream:
            for row in csv.DictReader(stream):
                yield json.loads(row['record_json'])
    finally:
        csv.field_size_limit(previous)
